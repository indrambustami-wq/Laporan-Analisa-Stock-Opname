import io
import re
from datetime import datetime

import pandas as pd
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)
from openpyxl.utils import get_column_letter


# ============================================================
# STREAMLIT CONFIG
# ============================================================

st.set_page_config(
    page_title="Laporan Analisa Stock Opname",
    page_icon="📦",
    layout="wide",
)


# ============================================================
# CONSTANTS
# ============================================================

NAVY = "1F497D"
YELLOW = "FFFF00"
WHITE = "FFFFFF"
BLACK = "000000"
RED = "FF0000"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_GRAY = "F2F2F2"

THIN_GRAY = Side(
    style="thin",
    color="B7B7B7",
)

DOUBLE_BLACK = Side(
    style="double",
    color=BLACK,
)

SHOPBAG_KEYWORDS = [
    "Shopping Bag",
    "Laundry",
    "Paper Bag",
    "Bag",
    "LB",
    "PBB",
    "SB",
]

CATEGORIES = [
    "KLOP",
    "SHOPBAG",
    "TERTUKAR",
    "MINUS",
    "PLUS",
]

REQUIRED_COLUMNS = [
    "OpnameNo",
    "Date",
    "Loc",
    "Golongan",
    "Artikel",
    "Barcode",
    "Qty POS",
    "Qty Fisik",
    "Price",
]

# Kolom yang tidak wajib ada di file input.
# Jika tidak ditemukan, akan otomatis diisi 0.
OPTIONAL_COLUMNS = [
    "HM",
]

# Alias nama kolom mentah -> nama kolom standar.
# Dipakai kalau nama kolom standar (mis. "HM") tidak
# ditemukan persis, tapi ada kolom lain yang maknanya sama.
COLUMN_ALIASES = {
    "HM": ["Cost"],
}

DETAIL_COLUMNS = [
    "Barcode",
    "Artikel",
    "Golongan",
    "Price",
    "Qty Fisik",
    "Qty POS",
    "Selisih",
    "Total",
    "HM",
    "Total HM",
    "Remark",
]


# ============================================================
# HEADER / COLUMN HELPERS
# ============================================================

def normalize_header(value):
    """Membersihkan nama header."""
    if pd.isna(value):
        return ""

    value = str(value).strip()
    value = re.sub(r"\s+", " ", value)

    return value


def canonical(value):
    """
    Menyamakan format nama kolom.

    Contoh:
        Qty POS
        QtyPOS
        qty_pos
        Qty-POS

    menjadi:
        qtypos
    """
    value = normalize_header(value).lower()

    return re.sub(
        r"[^a-z0-9]",
        "",
        value,
    )


def find_header_row(raw_df):
    """
    Deteksi otomatis baris header.
    Mencari hingga 50 baris pertama.
    """
    # Kolom opsional (mis. HM) dan alias-nya (mis. Cost)
    # tetap ikut dihitung skornya supaya baris header lebih
    # akurat terdeteksi kalau kolomnya ada, tapi tidak
    # menaikkan syarat minimal.
    alias_columns = [
        alias
        for aliases in COLUMN_ALIASES.values()
        for alias in aliases
    ]

    expected = {
        canonical(col)
        for col in REQUIRED_COLUMNS
        + OPTIONAL_COLUMNS
        + alias_columns
    }

    max_rows = min(
        len(raw_df),
        50,
    )

    best_row = None
    best_score = 0

    for row_idx in range(max_rows):
        values = [
            canonical(value)
            for value in raw_df.iloc[
                row_idx
            ].tolist()
        ]

        score = len(
            set(values) & expected
        )

        if score > best_score:
            best_score = score
            best_row = row_idx

    # Minimal kolom wajib yang harus ditemukan
    # (dikurangi toleransi 2 kolom hilang)
    minimum_required = max(
        1,
        len(REQUIRED_COLUMNS) - 2,
    )

    if best_score < minimum_required:
        return None

    return best_row


def map_standard_columns(df):
    """
    Mapping nama kolom mentah ke nama standar.
    Kalau nama kolom standar tidak ditemukan persis,
    coba cocokkan dengan alias-nya (lihat COLUMN_ALIASES).
    """
    df = df.copy()

    rename_map = {}

    for standard in REQUIRED_COLUMNS + OPTIONAL_COLUMNS:
        target = canonical(standard)

        match_col = None

        # 1. Cocokkan nama kolom standar persis
        for col in df.columns:
            if canonical(col) == target:
                match_col = col
                break

        # 2. Kalau tidak ketemu, coba alias-nya
        if match_col is None:
            for alias in COLUMN_ALIASES.get(
                standard,
                [],
            ):
                alias_target = canonical(alias)

                for col in df.columns:
                    if canonical(col) == alias_target:
                        match_col = col
                        break

                if match_col is not None:
                    break

        if match_col is not None:
            rename_map[match_col] = standard

    return df.rename(
        columns=rename_map
    )


# ============================================================
# FILE READING
# ============================================================

def read_csv_auto(file_bytes):
    """
    Membaca CSV dengan beberapa kemungkinan encoding
    dan mendeteksi header otomatis.
    """
    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            raw = pd.read_csv(
                io.BytesIO(file_bytes),
                header=None,
                encoding=encoding,
            )

            header_row = find_header_row(raw)

            if header_row is None:
                continue

            headers = [
                normalize_header(value)
                for value in raw.iloc[
                    header_row
                ].tolist()
            ]

            df = raw.iloc[
                header_row + 1:
            ].copy()

            df.columns = headers

            return df

        except Exception as exc:
            last_error = exc

    raise ValueError(
        f"CSV tidak dapat dibaca. Error: {last_error}"
    )


def load_excel_file(
    file_bytes,
    extension,
):
    """
    Membuka Excel berdasarkan extension.
    """
    if extension == ".xls":
        engine = "xlrd"
    else:
        engine = "openpyxl"

    return pd.ExcelFile(
        io.BytesIO(file_bytes),
        engine=engine,
    )


def load_excel_sheet(
    excel,
    sheet_name,
):
    """
    Membaca sheet Excel tanpa mengasumsikan
    header berada pada baris pertama.
    """
    raw = pd.read_excel(
        excel,
        sheet_name=sheet_name,
        header=None,
    )

    header_row = find_header_row(raw)

    if header_row is None:
        raise ValueError(
            f"Header tidak ditemukan pada sheet "
            f"'{sheet_name}'."
        )

    headers = [
        normalize_header(value)
        for value in raw.iloc[
            header_row
        ].tolist()
    ]

    df = raw.iloc[
        header_row + 1:
    ].copy()

    df.columns = headers

    # Buang kolom kosong / unnamed
    valid_columns = []

    for col in df.columns:
        text = str(col).strip()

        if (
            text != ""
            and not text.lower().startswith(
                "unnamed"
            )
        ):
            valid_columns.append(col)

    df = df[
        valid_columns
    ]

    return df


# ============================================================
# NUMBER PARSING
# ============================================================

def parse_number(value):
    """
    Konversi angka secara fleksibel.

    Contoh:
        1.234
        1,234
        1.234,56
        1,234.56
        Rp 10.000
        (10.000)
    """
    if pd.isna(value):
        return 0.0

    if isinstance(
        value,
        (int, float),
    ):
        return float(value)

    text = str(value).strip()

    if text == "":
        return 0.0

    text = (
        text
        .replace("Rp", "")
        .replace("rp", "")
        .replace(" ", "")
    )

    negative = (
        text.startswith("(")
        and text.endswith(")")
    )

    if negative:
        text = text[1:-1]

    # 1.234,56 atau 1,234.56
    if (
        "." in text
        and "," in text
    ):
        if text.rfind(",") > text.rfind("."):
            # Indonesia:
            # 1.234,56
            text = (
                text
                .replace(".", "")
                .replace(",", ".")
            )
        else:
            # International:
            # 1,234.56
            text = text.replace(
                ",",
                "",
            )

    # Hanya koma
    elif "," in text:
        parts = text.split(",")

        if (
            len(parts) > 1
            and len(parts[-1]) == 3
        ):
            # 1,234
            text = text.replace(
                ",",
                "",
            )
        else:
            # 1,25
            text = text.replace(
                ",",
                ".",
            )

    # Hanya titik
    elif "." in text:
        parts = text.split(".")

        if (
            len(parts) > 1
            and len(parts[-1]) == 3
        ):
            # 1.234
            text = text.replace(
                ".",
                "",
            )

    try:
        number = float(text)
    except Exception:
        number = 0.0

    if negative:
        number = -number

    return number


def clean_numeric(series):
    return series.map(
        parse_number
    ).fillna(0)


# ============================================================
# DATA CLEANING
# ============================================================

def prepare_data(df):
    """
    Cleaning data dan kalkulasi kolom turunan.
    """
    df = df.copy()

    # Hapus baris yang seluruhnya kosong
    df = df.dropna(
        how="all"
    ).reset_index(drop=True)

    # Standardisasi kolom
    df = map_standard_columns(
        df
    )

    missing = [
        col
        for col in REQUIRED_COLUMNS
        if col not in df.columns
    ]

    if missing:
        raise ValueError(
            "Kolom wajib tidak ditemukan: "
            + ", ".join(missing)
        )

    # Text columns
    text_columns = [
        "OpnameNo",
        "Date",
        "Loc",
        "Golongan",
        "Artikel",
        "Barcode",
    ]

    for col in text_columns:
        df[col] = (
            df[col]
            .fillna("")
            .astype(str)
            .str.strip()
        )

    # HM opsional: kalau tidak ada di file, isi 0
    if "HM" not in df.columns:
        df["HM"] = 0.0

    # Numeric columns
    for col in [
        "Qty POS",
        "Qty Fisik",
        "Price",
        "HM",
    ]:
        df[col] = clean_numeric(
            df[col]
        )

    # Remark optional
    if "Remark" not in df.columns:
        df["Remark"] = ""

    df["Remark"] = (
        df["Remark"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    # ========================================================
    # CALCULATIONS
    # ========================================================

    df["Selisih"] = (
        df["Qty Fisik"]
        - df["Qty POS"]
    )

    df["Total"] = (
        df["Selisih"]
        * df["Price"]
    )

    df["Total HM"] = (
        df["Selisih"]
        * df["HM"]
    )

    # Rounding
    df["Qty POS"] = df[
        "Qty POS"
    ].round(6)

    df["Qty Fisik"] = df[
        "Qty Fisik"
    ].round(6)

    df["Selisih"] = df[
        "Selisih"
    ].round(6)

    df["Price"] = df[
        "Price"
    ].round(2)

    df["HM"] = df[
        "HM"
    ].round(2)

    df["Total"] = df[
        "Total"
    ].round(2)

    df["Total HM"] = df[
        "Total HM"
    ].round(2)

    return df


# ============================================================
# BUSINESS LOGIC
# ============================================================

def is_shopbag(row):
    """
    SHOPBAG:
    Golongan atau Artikel mengandung salah satu keyword.
    """
    golongan = str(
        row["Golongan"]
    ).lower()

    artikel = str(
        row["Artikel"]
    ).lower()

    for keyword in SHOPBAG_KEYWORDS:
        keyword = keyword.lower()

        if (
            keyword in golongan
            or keyword in artikel
        ):
            return True

    return False


def categorize(df):
    """
    Prioritas:

    1. SHOPBAG
    2. TERTUKAR
    3. MINUS
    4. PLUS
    5. KLOP
    """
    df = df.copy()

    df["Kategori"] = ""

    # ========================================================
    # 1. SHOPBAG
    # ========================================================

    shopbag_mask = df.apply(
        is_shopbag,
        axis=1,
    )

    df.loc[
        shopbag_mask,
        "Kategori",
    ] = "SHOPBAG"

    # ========================================================
    # 2. TERTUKAR
    # ========================================================

    remaining = (
        df["Kategori"] == ""
    )

    article_signs = (
        df.loc[remaining]
        .groupby("Artikel")[
            "Selisih"
        ]
        .agg(
            has_minus=lambda x: (
                x < 0
            ).any(),
            has_plus=lambda x: (
                x > 0
            ).any(),
        )
    )

    exchanged_articles = (
        article_signs[
            article_signs[
                "has_minus"
            ]
            & article_signs[
                "has_plus"
            ]
        ].index
    )

    tertukar_mask = (
        (df["Kategori"] == "")
        & df["Artikel"].isin(
            exchanged_articles
        )
        & (df["Selisih"] != 0)
    )

    df.loc[
        tertukar_mask,
        "Kategori",
    ] = "TERTUKAR"

    # ========================================================
    # 3. MINUS
    # ========================================================

    minus_mask = (
        (df["Kategori"] == "")
        & (df["Selisih"] < 0)
    )

    df.loc[
        minus_mask,
        "Kategori",
    ] = "MINUS"

    # ========================================================
    # 4. PLUS
    # ========================================================

    plus_mask = (
        (df["Kategori"] == "")
        & (df["Selisih"] > 0)
    )

    df.loc[
        plus_mask,
        "Kategori",
    ] = "PLUS"

    # ========================================================
    # 5. KLOP
    # ========================================================

    klop_mask = (
        (df["Kategori"] == "")
        & (df["Selisih"] == 0)
    )

    df.loc[
        klop_mask,
        "Kategori",
    ] = "KLOP"

    # ========================================================
    # REMARK
    # ========================================================
    # Default: nama kategori (Klop/Tertukar/Minus/Plus).
    # Khusus SHOPBAG, ikuti tanda Selisih (Minus/Plus/Klop)
    # karena SHOPBAG sendiri bisa berisi selisih apa saja.

    def get_remark(row):
        category = row["Kategori"]
        selisih = row["Selisih"]

        if category == "SHOPBAG":
            if selisih < 0:
                return "Minus"

            if selisih > 0:
                return "Plus"

            return "Klop"

        return category.capitalize()

    df["Remark"] = df.apply(
        get_remark,
        axis=1,
    )

    return df


# ============================================================
# SUMMARY
# ============================================================

def create_summary(df):
    rows = []

    for category in CATEGORIES:
        part = df[
            df["Kategori"]
            == category
        ]

        rows.append(
            {
                "Kategori": category,
                "Qty Fisik": part[
                    "Qty Fisik"
                ].sum(),
                "Qty POS": part[
                    "Qty POS"
                ].sum(),
                "Selisih": part[
                    "Selisih"
                ].sum(),
                "Nilai Selisih (Rp)": part[
                    "Total"
                ].sum(),
                "HM": part[
                    "HM"
                ].sum(),
                "Total HM": part[
                    "Total HM"
                ].sum(),
            }
        )

    # Grand Total
    rows.append(
        {
            "Kategori": "GRAND TOTAL",
            "Qty Fisik": df[
                "Qty Fisik"
            ].sum(),
            "Qty POS": df[
                "Qty POS"
            ].sum(),
            "Selisih": df[
                "Selisih"
            ].sum(),
            "Nilai Selisih (Rp)": df[
                "Total"
            ].sum(),
            "HM": df[
                "HM"
            ].sum(),
            "Total HM": df[
                "Total HM"
            ].sum(),
        }
    )

    return pd.DataFrame(
        rows
    )


def get_metadata(df):
    def first_value(column):
        values = (
            df[column]
            .dropna()
            .astype(str)
            .str.strip()
        )

        values = values[
            values != ""
        ]

        if len(values):
            return values.iloc[0]

        return ""

    return {
        "Loc": first_value("Loc"),
        "Date": first_value("Date"),
        "OpnameNo": first_value(
            "OpnameNo"
        ),
    }


# ============================================================
# EXCEL STYLE FUNCTIONS
# ============================================================

def apply_number_format(
    cell,
    money=False,
    value=None,
):
    """
    Format angka konsisten dengan section
    NOTE / PERHITUNGAN DEDUCT:

    - money=True  -> pakai prefix "Rp"
                      (mis. -Rp26.404 / Rp0 / Rp500.000)
    - money=False -> angka polos
                      (mis. -12 / 0 / 17)

    Nilai negatif otomatis ditampilkan warna merah.

    Parameter `value` dipakai untuk menentukan warna merah
    kalau cell.value berupa rumus (string, mis. "=E2-F2")
    sehingga tandanya tidak bisa dibaca langsung dari
    cell.value. Kalau tidak diisi, dipakai cell.value.
    """

    if money:
        cell.number_format = (
            '"Rp"#,##0;-"Rp"#,##0;"Rp"0'
        )
    else:
        cell.number_format = "#,##0"

    check_value = (
        cell.value
        if value is None
        else value
    )

    if (
        isinstance(
            check_value,
            (int, float),
        )
        and check_value < 0
    ):
        old_font = cell.font

        cell.font = Font(
            name=old_font.name,
            size=old_font.size,
            bold=old_font.bold,
            italic=old_font.italic,
            color=RED,
        )


def apply_border(
    cell,
    left=None,
    right=None,
    top=None,
    bottom=None,
):
    """
    Mempertahankan border lain ketika
    mengubah satu sisi border.
    """
    old = cell.border

    cell.border = Border(
        left=left or old.left,
        right=right or old.right,
        top=top or old.top,
        bottom=bottom or old.bottom,
    )


def style_summary_header(
    ws,
    row,
    start_col,
):
    headers = [
        "Kategori",
        "Qty Fisik",
        "Qty POS",
        "Selisih",
        "Nilai Selisih (Rp)",
        "HM",
        "Total HM",
    ]

    for offset, header in enumerate(
        headers
    ):
        cell = ws.cell(
            row=row,
            column=start_col + offset,
            value=header,
        )

        is_hm = header in [
            "HM",
            "Total HM",
        ]

        cell.fill = PatternFill(
            "solid",
            fgColor=(
                YELLOW
                if is_hm
                else NAVY
            ),
        )

        cell.font = Font(
            bold=True,
            color=(
                BLACK
                if is_hm
                else WHITE
            ),
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )


def style_detail_header(
    ws,
    row,
):
    for col_idx, header in enumerate(
        DETAIL_COLUMNS,
        start=1,
    ):
        cell = ws.cell(
            row=row,
            column=col_idx,
            value=header,
        )

        is_hm = header in [
            "HM",
            "Total HM",
        ]

        cell.fill = PatternFill(
            "solid",
            fgColor=(
                YELLOW
                if is_hm
                else NAVY
            ),
        )

        cell.font = Font(
            bold=True,
            color=(
                BLACK
                if is_hm
                else WHITE
            ),
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )


# ============================================================
# CREATE FINAL EXCEL
# ============================================================

def create_final_excel(
    df,
    metadata,
    summary,
    penalty_percent,
    compensation,
    compensation_note="",
):
    """
    Membuat workbook FINAL.
    """

    wb = Workbook()

    ws = wb.active
    ws.title = "FINAL"

    ws.sheet_view.showGridLines = False

    # ========================================================
    # HEADER METADATA A-C
    # ========================================================

    ws.merge_cells("A1:C1")

    ws["A1"] = (
        "LAPORAN ANALISA FINAL"
    )

    ws["A1"].font = Font(
        bold=True,
        size=15,
        color=NAVY,
    )

    ws["A1"].alignment = Alignment(
        vertical="center"
    )

    metadata_fields = [
        ("Lokasi / Loc", metadata["Loc"]),
        ("Tanggal SO", metadata["Date"]),
        ("No. Opname", metadata["OpnameNo"]),
    ]

    for offset, (label, value) in enumerate(metadata_fields):
        row = 2 + offset

        label_cell = ws.cell(
            row=row,
            column=1,
            value=label,
        )

        label_cell.font = Font(
            bold=True,
        )

        label_cell.alignment = Alignment(
            vertical="center"
        )

        label_cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_GRAY,
        )

        # Kolom C ikut diberi warna sebelum di-merge
        # dengan kolom B (nilai)
        ws.cell(
            row=row,
            column=3,
        ).fill = PatternFill(
            "solid",
            fgColor=LIGHT_GRAY,
        )

        # Value dipisah ke kolom B:C (merge biar muat)
        ws.merge_cells(
            start_row=row,
            start_column=2,
            end_row=row,
            end_column=3,
        )

        value_cell = ws.cell(
            row=row,
            column=2,
            value=value,
        )

        value_cell.alignment = Alignment(
            vertical="center"
        )

        value_cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_GRAY,
        )

    # ========================================================
    # SUMMARY TABLE D-J
    # ========================================================

    summary_start_col = 4
    summary_header_row = 1

    style_summary_header(
        ws,
        summary_header_row,
        summary_start_col,
    )

    # BUG FIX:
    # Tidak menggunakan itertuples() / _1 / _2 / _3.
    # Menggunakan iterrows() dengan nama kolom eksplisit.
    for excel_row, (
        _,
        record,
    ) in enumerate(
        summary.iterrows(),
        start=2,
    ):
        values = [
            record["Kategori"],
            record["Qty Fisik"],
            record["Qty POS"],
            record["Selisih"],
            record["Nilai Selisih (Rp)"],
            record["HM"],
            record["Total HM"],
        ]

        for offset, value in enumerate(
            values
        ):
            cell = ws.cell(
                row=excel_row,
                column=summary_start_col
                + offset,
                value=value,
            )

            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )

            if offset == 0:
                cell.font = Font(
                    bold=True
                )

            if offset >= 1:
                apply_number_format(
                    cell,
                    money=offset in [
                        4,
                        5,
                        6,
                    ],
                )

        # Grand Total
        if (
            record["Kategori"]
            == "GRAND TOTAL"
        ):
            grand_values = [
                None,
                record["Qty Fisik"],
                record["Qty POS"],
                record["Selisih"],
                record["Nilai Selisih (Rp)"],
                record["HM"],
                record["Total HM"],
            ]

            for offset in range(7):
                cell = ws.cell(
                    row=excel_row,
                    column=summary_start_col
                    + offset,
                )

                raw_value = grand_values[
                    offset
                ]

                is_negative = (
                    isinstance(
                        raw_value,
                        (int, float),
                    )
                    and raw_value < 0
                )

                cell.font = Font(
                    bold=True,
                    color=(
                        RED
                        if is_negative
                        else None
                    ),
                )

                apply_border(
                    cell,
                    bottom=DOUBLE_BLACK,
                )

    # ========================================================
    # DETAIL TABLE
    # ========================================================

    # Huruf kolom Excel untuk setiap nama kolom di
    # DETAIL_COLUMNS, dipakai untuk menulis rumus
    # (Selisih, Total, Total HM).
    col_letter = {
        name: get_column_letter(idx)
        for idx, name in enumerate(
            DETAIL_COLUMNS,
            start=1,
        )
    }

    current_row = 9

    for category in CATEGORIES:

        part = df[
            df["Kategori"]
            == category
        ].copy()

        # ----------------------------------------------------
        # CATEGORY TITLE
        # ----------------------------------------------------

        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=len(
                DETAIL_COLUMNS
            ),
        )

        title_cell = ws.cell(
            row=current_row,
            column=1,
            value=(
                f"KATEGORI: {category}"
            ),
        )

        title_cell.font = Font(
            bold=True,
            size=12,
            color=NAVY,
        )

        title_cell.fill = PatternFill(
            "solid",
            fgColor=LIGHT_BLUE,
        )

        title_cell.alignment = Alignment(
            vertical="center"
        )

        current_row += 1

        # ----------------------------------------------------
        # DETAIL HEADER
        # ----------------------------------------------------

        style_detail_header(
            ws,
            current_row,
        )

        current_row += 1

        # ----------------------------------------------------
        # DETAIL DATA
        # ----------------------------------------------------

        data_start_row = current_row

        for _, source in part.iterrows():

            for col_idx, column in enumerate(
                DETAIL_COLUMNS,
                start=1,
            ):
                value = source.get(
                    column,
                    "",
                )

                # Jangan tulis NaN
                if pd.isna(value):
                    value = ""

                # Kolom hasil kalkulasi ditulis sebagai
                # rumus Excel (bukan angka statis) supaya
                # bisa diaudit langsung dari file.
                formula = None

                if column == "Selisih":
                    formula = (
                        f"={col_letter['Qty Fisik']}"
                        f"{current_row}-"
                        f"{col_letter['Qty POS']}"
                        f"{current_row}"
                    )

                elif column == "Total":
                    formula = (
                        f"={col_letter['Selisih']}"
                        f"{current_row}*"
                        f"{col_letter['Price']}"
                        f"{current_row}"
                    )

                elif column == "Total HM":
                    formula = (
                        f"={col_letter['Selisih']}"
                        f"{current_row}*"
                        f"{col_letter['HM']}"
                        f"{current_row}"
                    )

                cell = ws.cell(
                    row=current_row,
                    column=col_idx,
                    value=(
                        formula
                        if formula
                        is not None
                        else value
                    ),
                )

                cell.border = Border(
                    left=THIN_GRAY,
                    right=THIN_GRAY,
                    top=THIN_GRAY,
                    bottom=THIN_GRAY,
                )

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=False,
                )

                if column in [
                    "Price",
                    "Qty Fisik",
                    "Qty POS",
                    "Selisih",
                    "Total",
                    "HM",
                    "Total HM",
                ]:
                    apply_number_format(
                        cell,
                        money=column in [
                            "Price",
                            "Total",
                            "HM",
                            "Total HM",
                        ],
                        # Sel Selisih/Total/Total HM
                        # berisi rumus (string), jadi warna
                        # merah ditentukan dari nilai asli
                        # yang sudah dihitung Python.
                        value=value,
                    )

            current_row += 1

        # ----------------------------------------------------
        # SUBTOTAL
        # ----------------------------------------------------

        subtotal_row = current_row

        ws.cell(
            row=subtotal_row,
            column=1,
            value="SUBTOTAL",
        )

        subtotal_map = {
            5: "Qty Fisik",
            6: "Qty POS",
            7: "Selisih",
            8: "Total",
            10: "Total HM",
        }

        # Style semua cell subtotal
        for col_idx in range(
            1,
            len(DETAIL_COLUMNS) + 1,
        ):
            cell = ws.cell(
                row=subtotal_row,
                column=col_idx,
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=LIGHT_GREEN,
            )

            cell.font = Font(
                bold=True
            )

            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=DOUBLE_BLACK,
            )

        # Isi subtotal
        has_data_rows = (
            len(part) > 0
        )

        sum_start = data_start_row
        sum_end = current_row - 1

        for col_idx, source_col in subtotal_map.items():

            numeric_value = part[
                source_col
            ].sum()

            letter = col_letter[
                source_col
            ]

            if has_data_rows:
                formula = (
                    f"=SUM({letter}"
                    f"{sum_start}:"
                    f"{letter}"
                    f"{sum_end})"
                )
            else:
                # Tidak ada baris data,
                # tulis 0 langsung
                formula = 0

            cell = ws.cell(
                row=subtotal_row,
                column=col_idx,
                value=formula,
            )

            apply_number_format(
                cell,
                money=source_col in [
                    "Total",
                    "Total HM",
                ],
                # Rumus SUM tidak bisa dibaca
                # tandanya langsung, jadi pakai
                # nilai asli hasil hitung Python.
                value=numeric_value,
            )

        current_row += 2

    # ========================================================
    # NOTE / DEDUCT
    # ========================================================

    def apply_rupiah_format(cell):
        """
        Format khusus untuk section NOTE/PERHITUNGAN DEDUCT:
        -Rp26.404 / Rp0 / Rp500.000
        """
        cell.number_format = (
            '"Rp"#,##0;-"Rp"#,##0;"Rp"0'
        )

    note_start = current_row + 1

    ws.merge_cells(
        start_row=note_start,
        start_column=1,
        end_row=note_start,
        end_column=3,
    )

    ws.cell(
        row=note_start,
        column=1,
        value=(
            "NOTE / PERHITUNGAN DEDUCT"
        ),
    )

    ws.cell(
        row=note_start,
        column=1,
    ).font = Font(
        bold=True,
        size=12,
        color=NAVY,
    )

    note_row = note_start + 1

    # --------------------------------------------------------
    # NOTE HEADER
    # --------------------------------------------------------

    note_headers = [
        "Keterangan",
        "Qty Selisih",
        "Total HM",
    ]

    for col_idx, header in enumerate(
        note_headers,
        start=1,
    ):
        cell = ws.cell(
            row=note_row,
            column=col_idx,
            value=header,
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=NAVY,
        )

        cell.font = Font(
            bold=True,
            color=WHITE,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )

    note_row += 1

    # --------------------------------------------------------
    # CATEGORY VALUES
    # --------------------------------------------------------

    def category_values(category):
        part = df[
            df["Kategori"]
            == category
        ]

        return (
            part["Selisih"].sum(),
            part["Total HM"].sum(),
        )

    shopbag_qty, shopbag_hm = (
        category_values(
            "SHOPBAG"
        )
    )

    minus_qty, minus_hm = (
        category_values(
            "MINUS"
        )
    )

    plus_qty, plus_hm = (
        category_values(
            "PLUS"
        )
    )

    note_items = [
        (
            "Shopbag",
            shopbag_qty,
            shopbag_hm,
        ),
        (
            "Minus",
            minus_qty,
            minus_hm,
        ),
        (
            "Plus",
            plus_qty,
            plus_hm,
        ),
    ]

    for label, qty, hm in note_items:

        values = [
            label,
            qty,
            hm,
        ]

        for col_idx, value in enumerate(
            values,
            start=1,
        ):
            cell = ws.cell(
                row=note_row,
                column=col_idx,
                value=value,
            )

            cell.border = Border(
                left=THIN_GRAY,
                right=THIN_GRAY,
                top=THIN_GRAY,
                bottom=THIN_GRAY,
            )

        qty_cell = ws.cell(
            row=note_row,
            column=2,
        )

        qty_cell.number_format = (
            "#,##0"
        )

        hm_cell = ws.cell(
            row=note_row,
            column=3,
        )

        apply_rupiah_format(
            hm_cell
        )

        # Nilai negatif ditandai merah
        if qty < 0:
            qty_cell.font = Font(
                color=RED
            )

        if hm < 0:
            hm_cell.font = Font(
                color=RED
            )

        note_row += 1

    # --------------------------------------------------------
    # TOTAL SELISIH
    # --------------------------------------------------------

    total_qty = (
        shopbag_qty
        + minus_qty
        + plus_qty
    )

    total_hm = (
        shopbag_hm
        + minus_hm
        + plus_hm
    )

    total_row = note_row

    total_values = [
        "Total Selisih",
        total_qty,
        total_hm,
    ]

    for col_idx, value in enumerate(
        total_values,
        start=1,
    ):
        cell = ws.cell(
            row=total_row,
            column=col_idx,
            value=value,
        )

        cell.font = Font(
            bold=True
        )

        cell.border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=DOUBLE_BLACK,
            bottom=DOUBLE_BLACK,
        )

    ws.cell(
        row=total_row,
        column=2,
    ).number_format = "#,##0"

    apply_rupiah_format(
        ws.cell(
            row=total_row,
            column=3,
        )
    )

    note_row += 1

    # Baris kosong (spasi)
    note_row += 1

    # --------------------------------------------------------
    # TOTAL HM + DENDA
    # --------------------------------------------------------

    denda_value = (
        total_hm
        * penalty_percent
    )

    total_hm_denda = (
        total_hm
        + denda_value
    )

    denda_row = note_row

    ws.cell(
        row=denda_row,
        column=1,
        value="Total HM + Denda",
    )

    denda_percent_cell = ws.cell(
        row=denda_row,
        column=2,
        value=penalty_percent,
    )

    denda_percent_cell.number_format = (
        "0%"
    )

    denda_value_cell = ws.cell(
        row=denda_row,
        column=3,
        value=total_hm_denda,
    )

    apply_rupiah_format(
        denda_value_cell
    )

    for col_idx in range(1, 4):
        ws.cell(
            row=denda_row,
            column=col_idx,
        ).border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )

    if total_hm_denda < 0:
        denda_value_cell.font = Font(
            color=RED,
        )

    note_row += 1

    # --------------------------------------------------------
    # KOMPENSASI
    # --------------------------------------------------------

    kompensasi_row = note_row

    ws.cell(
        row=kompensasi_row,
        column=1,
        value="Kompensasi",
    )

    kompensasi_cell = ws.cell(
        row=kompensasi_row,
        column=3,
        value=compensation,
    )

    apply_rupiah_format(
        kompensasi_cell
    )

    for col_idx in range(1, 4):
        ws.cell(
            row=kompensasi_row,
            column=col_idx,
        ).border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=THIN_GRAY,
            bottom=THIN_GRAY,
        )

    note_row += 1

    # Baris kosong (spasi)
    note_row += 1

    # --------------------------------------------------------
    # TOTAL DEDUCT
    # --------------------------------------------------------

    total_deduct = max(
        0,
        abs(total_hm_denda)
        - compensation,
    )

    deduct_row = note_row

    ws.cell(
        row=deduct_row,
        column=1,
        value="Total Deduct",
    ).font = Font(
        bold=True
    )

    for col_idx in range(1, 4):
        ws.cell(
            row=deduct_row,
            column=col_idx,
        ).border = Border(
            left=THIN_GRAY,
            right=THIN_GRAY,
            top=DOUBLE_BLACK,
            bottom=DOUBLE_BLACK,
        )

    # Deduct wajib kuning
    deduct_cell = ws.cell(
        row=deduct_row,
        column=3,
        value=total_deduct,
    )

    deduct_cell.fill = PatternFill(
        "solid",
        fgColor=YELLOW,
    )

    deduct_cell.font = Font(
        bold=True,
        color=BLACK,
    )

    apply_rupiah_format(
        deduct_cell
    )

    note_row += 1

    # --------------------------------------------------------
    # CATATAN KOMPENSASI (teks miring di bawah tabel)
    # --------------------------------------------------------

    if compensation_note:

        catatan_row = note_row + 1

        ws.merge_cells(
            start_row=catatan_row,
            start_column=1,
            end_row=catatan_row,
            end_column=3,
        )

        catatan_cell = ws.cell(
            row=catatan_row,
            column=1,
            value=(
                f"Catatan Kompensasi: "
                f"{compensation_note}"
            ),
        )

        catatan_cell.font = Font(
            italic=True,
        )

        catatan_cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

        note_row = catatan_row

    # ========================================================
    # GLOBAL EXCEL FORMATTING
    # ========================================================

    ws.freeze_panes = "A11"

    # --------------------------------------------------------
    # Auto-fit lebar kolom (dihitung dari isi tiap kolom)
    # --------------------------------------------------------

    def estimate_formatted_length(
        raw_value,
        money=False,
    ):
        """
        Perkiraan panjang teks angka setelah diformat
        (dipakai untuk kolom yang isinya rumus, karena
        openpyxl tidak bisa membaca hasil rumus).
        """
        try:
            number = float(raw_value)
        except (TypeError, ValueError):
            return len(
                str(raw_value)
            )

        text = f"{abs(number):,.0f}"

        if money:
            text = "Rp" + text

        if number < 0:
            text = "-" + text

        return len(text)

    MIN_COLUMN_WIDTH = 10
    MAX_COLUMN_WIDTH = 40
    COLUMN_PADDING = 2

    # Sel yang jadi "anchor" dari merge lebih dari 1 kolom
    # dilewati dari perhitungan (teksnya tidak perlu muat
    # dalam satu kolom saja, mis. judul & catatan panjang).
    wide_merge_anchors = set()

    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.max_col
            > merged_range.min_col
        ):
            wide_merge_anchors.add(
                (
                    merged_range.min_row,
                    merged_range.min_col,
                )
            )

    max_lengths = {}

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue

            if (
                cell.row,
                cell.column,
            ) in wide_merge_anchors:
                continue

            text = str(cell.value)

            # Sel berisi rumus: panjang teks rumusnya
            # tidak merepresentasikan tampilan akhir,
            # jadi dilewati (diestimasi terpisah di bawah).
            if text.startswith("="):
                continue

            col = cell.column_letter

            max_lengths[col] = max(
                max_lengths.get(
                    col, 0
                ),
                len(text),
            )

    # Kolom hasil rumus di tabel detail (Selisih, Total,
    # Total HM) diestimasi dari nilai asli di dataframe.
    formula_columns = {
        col_letter["Selisih"]: (
            df["Selisih"],
            False,
        ),
        col_letter["Total"]: (
            df["Total"],
            True,
        ),
        col_letter["Total HM"]: (
            df["Total HM"],
            True,
        ),
    }

    for col, (
        series,
        money,
    ) in formula_columns.items():

        if len(series) == 0:
            continue

        estimated = max(
            estimate_formatted_length(
                v,
                money=money,
            )
            for v in series
        )

        max_lengths[col] = max(
            max_lengths.get(
                col, 0
            ),
            estimated,
        )

    for col, length in max_lengths.items():
        width = min(
            max(
                length
                + COLUMN_PADDING,
                MIN_COLUMN_WIDTH,
            ),
            MAX_COLUMN_WIDTH,
        )

        ws.column_dimensions[
            col
        ].width = width

    # --------------------------------------------------------
    # Row height
    # --------------------------------------------------------

    ws.row_dimensions[1].height = 26

    # --------------------------------------------------------
    # General alignment
    # --------------------------------------------------------

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                # Jangan mengubah alignment note yang sudah wrap
                if cell.alignment.wrap_text is None:
                    cell.alignment = Alignment(
                        vertical="center"
                    )

    # ========================================================
    # OUTPUT BYTES
    # ========================================================

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# SIDEBAR
# ============================================================

st.title(
    "📦 Laporan Analisa Stock Opname"
)

st.caption(
    "Upload file DATA mentah untuk menghasilkan "
    "analisa otomatis dan sheet FINAL."
)

with st.sidebar:

    st.header(
        "⚙️ Parameter"
    )

    uploaded_file = st.file_uploader(
        "Input File",
        type=[
            "xlsx",
            "xls",
            "csv",
        ],
        help=(
            "Mendukung Excel XLSX/XLS dan CSV."
        ),
    )

    penalty_percent = st.number_input(
        "Persentase Denda (%)",
        min_value=0.0,
        max_value=100.0,
        value=30.0,
        step=1.0,
        format="%.2f",
    )

    compensation = st.number_input(
        "Nilai Kompensasi (Rp)",
        min_value=0.0,
        value=500_000.0,
        step=50_000.0,
        format="%.0f",
    )

    compensation_note = st.text_area(
        "Catatan Kompensasi",
        placeholder=(
            "Contoh: Kompensasi sesuai "
            "kebijakan periode berjalan."
        ),
    )

    st.markdown("---")

    st.caption(
        "Urutan kategorisasi:"
    )

    st.caption(
        "SHOPBAG → TERTUKAR → MINUS → PLUS → KLOP"
    )


# ============================================================
# NO FILE
# ============================================================

if uploaded_file is None:

    st.info(
        "Silakan upload file melalui sidebar."
    )

    st.markdown(
        """
### Kolom yang diperlukan

File mentah minimal harus memiliki:

- `OpnameNo`
- `Date`
- `Loc`
- `Golongan`
- `Artikel`
- `Barcode`
- `Qty POS`
- `Qty Fisik`
- `Price`

Kolom berikut **opsional** (kalau tidak ada, otomatis diisi 0):

- `HM`

Header dapat berada di bawah beberapa baris awal karena
aplikasi melakukan deteksi header secara otomatis.
"""
    )

    st.stop()


# ============================================================
# READ INPUT
# ============================================================

file_bytes = uploaded_file.getvalue()
filename = uploaded_file.name.lower()

try:

    if filename.endswith(".csv"):

        df_raw = read_csv_auto(
            file_bytes
        )

        selected_sheet = "CSV"

    else:

        if (
            filename.endswith(".xls")
            and not filename.endswith(
                ".xlsx"
            )
        ):
            extension = ".xls"
        else:
            extension = ".xlsx"

        excel = load_excel_file(
            file_bytes,
            extension,
        )

        sheets = excel.sheet_names

        # Prioritas DATA
        if "DATA" in sheets:
            default_index = sheets.index(
                "DATA"
            )

        elif "Data" in sheets:
            default_index = sheets.index(
                "Data"
            )

        else:
            default_index = 0

        selected_sheet = st.sidebar.selectbox(
            "Sheet Data",
            sheets,
            index=default_index,
        )

        df_raw = load_excel_sheet(
            excel,
            selected_sheet,
        )

except Exception as exc:

    st.error(
        f"Gagal membaca file: {exc}"
    )

    st.stop()


# ============================================================
# PROCESS
# ============================================================

try:

    hm_missing = (
        "HM"
        not in map_standard_columns(
            df_raw
        ).columns
    )

    df = prepare_data(
        df_raw
    )

    df = categorize(
        df
    )

    summary = create_summary(
        df
    )

    metadata = get_metadata(
        df
    )

    if hm_missing:
        st.warning(
            "⚠️ Kolom **HM** tidak ditemukan di file "
            "input. Nilai HM otomatis diisi 0 untuk "
            "seluruh baris."
        )

except Exception as exc:

    st.error(
        f"Gagal memproses data: {exc}"
    )

    with st.expander(
        "🔎 Header yang Terdeteksi"
    ):
        st.write(
            list(df_raw.columns)
        )

    st.stop()


# ============================================================
# DEDUCT CALCULATION
# ============================================================

total_hm = df[
    "Total HM"
].sum()

denda_value = (
    total_hm
    * penalty_percent
    / 100
)

total_hm_denda = (
    total_hm
    + denda_value
)

total_deduct = max(
    0,
    abs(total_hm_denda)
    - compensation,
)


# ============================================================
# METADATA
# ============================================================

st.markdown(
    "## 📌 Informasi Stock Opname"
)

m1, m2, m3 = st.columns(3)

with m1:
    st.metric(
        "Lokasi / Loc",
        metadata["Loc"]
        or "-",
    )

with m2:
    st.metric(
        "Tanggal SO",
        metadata["Date"]
        or "-",
    )

with m3:
    st.metric(
        "No. Opname",
        metadata["OpnameNo"]
        or "-",
    )


# ============================================================
# KPI
# ============================================================

st.markdown(
    "## 📊 KPI"
)

grand = summary[
    summary["Kategori"]
    == "GRAND TOTAL"
].iloc[0]

k1, k2, k3, k4, k5 = st.columns(5)

with k1:
    st.metric(
        "Total Qty Fisik",
        f"{grand['Qty Fisik']:,.0f}",
    )

with k2:
    st.metric(
        "Total Qty POS",
        f"{grand['Qty POS']:,.0f}",
    )

with k3:
    st.metric(
        "Total Selisih Qty",
        f"{grand['Selisih']:,.0f}",
    )

with k4:
    st.metric(
        "Total Nilai Selisih",
        f"Rp {grand['Nilai Selisih (Rp)']:,.0f}",
    )

with k5:
    st.metric(
        "Total Deduct",
        f"Rp {total_deduct:,.0f}",
    )


# ============================================================
# EXTRA KPI
# ============================================================

a1, a2 = st.columns(2)

with a1:
    st.metric(
        f"Total HM + Denda ({penalty_percent:.0f}%)",
        f"Rp {total_hm_denda:,.0f}",
    )

with a2:
    st.metric(
        "Kompensasi",
        f"Rp {compensation:,.0f}",
    )


# ============================================================
# SUMMARY
# ============================================================

st.markdown(
    "## 📋 Rekapitulasi Kategori"
)

summary_display = summary.copy()

for col in [
    "Qty Fisik",
    "Qty POS",
    "Selisih",
    "Nilai Selisih (Rp)",
    "HM",
    "Total HM",
]:

    summary_display[col] = (
        summary_display[col]
        .map(
            lambda value:
            f"{value:,.0f}"
        )
    )

st.dataframe(
    summary_display,
    use_container_width=True,
    hide_index=True,
)


# ============================================================
# CATEGORY TABS
# ============================================================

st.markdown(
    "## 🔍 Detail Data Per Kategori"
)

tabs = st.tabs(
    [
        (
            f"{category} "
            f"({len(df[df['Kategori'] == category])})"
        )
        for category in CATEGORIES
    ]
)

for tab, category in zip(
    tabs,
    CATEGORIES,
):

    with tab:

        part = df[
            df["Kategori"]
            == category
        ].copy()

        display_columns = [
            "Barcode",
            "Artikel",
            "Golongan",
            "Price",
            "Qty Fisik",
            "Qty POS",
            "Selisih",
            "Total",
            "HM",
            "Total HM",
            "Remark",
        ]

        st.dataframe(
            part[
                display_columns
            ],
            use_container_width=True,
            hide_index=True,
            column_config={
                "Price":
                    st.column_config.NumberColumn(
                        "Price",
                        format="%,.0f",
                    ),

                "Qty Fisik":
                    st.column_config.NumberColumn(
                        "Qty Fisik",
                        format="%,.0f",
                    ),

                "Qty POS":
                    st.column_config.NumberColumn(
                        "Qty POS",
                        format="%,.0f",
                    ),

                "Selisih":
                    st.column_config.NumberColumn(
                        "Selisih",
                        format="%,.0f",
                    ),

                "Total":
                    st.column_config.NumberColumn(
                        "Total",
                        format="Rp %,.0f",
                    ),

                "HM":
                    st.column_config.NumberColumn(
                        "HM",
                        format="%,.0f",
                    ),

                "Total HM":
                    st.column_config.NumberColumn(
                        "Total HM",
                        format="Rp %,.0f",
                    ),
            },
        )

        if len(part) > 0:

            c1, c2, c3 = st.columns(3)

            with c1:
                st.metric(
                    "Qty Selisih",
                    f"{part['Selisih'].sum():,.0f}",
                )

            with c2:
                st.metric(
                    "Nilai Selisih",
                    f"Rp {part['Total'].sum():,.0f}",
                )

            with c3:
                st.metric(
                    "Total HM",
                    f"Rp {part['Total HM'].sum():,.0f}",
                )


# ============================================================
# DEDUCT DETAIL
# ============================================================

st.markdown(
    "## 💰 Detail Perhitungan Deduct"
)

d1, d2, d3, d4 = st.columns(4)

with d1:
    st.metric(
        "Total HM",
        f"Rp {total_hm:,.0f}",
    )

with d2:
    st.metric(
        f"Denda {penalty_percent:.0f}%",
        f"Rp {denda_value:,.0f}",
    )

with d3:
    st.metric(
        "Kompensasi",
        f"Rp {compensation:,.0f}",
    )

with d4:
    st.metric(
        "Total Deduct",
        f"Rp {total_deduct:,.0f}",
    )

if compensation_note:
    st.info(
        f"Catatan Kompensasi: "
        f"{compensation_note}"
    )


# ============================================================
# DATA PREVIEW
# ============================================================

with st.expander(
    "🔎 Preview Data Setelah Cleaning & Kategorisasi"
):

    preview_columns = [
        "OpnameNo",
        "Date",
        "Loc",
        "Golongan",
        "Artikel",
        "Barcode",
        "Qty POS",
        "Qty Fisik",
        "Price",
        "HM",
        "Selisih",
        "Total",
        "Total HM",
        "Kategori",
        "Remark",
    ]

    st.dataframe(
        df[
            preview_columns
        ],
        use_container_width=True,
        hide_index=True,
    )


# ============================================================
# EXPORT
# ============================================================

st.markdown("---")

st.markdown(
    "## 📥 Download Laporan FINAL"
)

try:

    excel_bytes = create_final_excel(
        df=df,
        metadata=metadata,
        summary=summary,
        penalty_percent=(
            penalty_percent
            / 100
        ),
        compensation=compensation,
        compensation_note=(
            compensation_note
        ),
    )

    safe_loc = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        metadata["Loc"]
        or "ALL",
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    output_filename = (
        "FINAL_STOCK_OPNAME_"
        f"{safe_loc}_"
        f"{timestamp}.xlsx"
    )

    st.download_button(
        label=(
            "📥 Download Laporan FINAL (.xlsx)"
        ),
        data=excel_bytes,
        file_name=output_filename,
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
        type="primary",
    )

except Exception as exc:

    st.error(
        f"Gagal membuat Excel FINAL: {exc}"
    )
